#!/usr/bin/env python3
"""
Renomeia arquivos ANBIMA existentes para o padrão organizado:
  debentures/debenture-DD-MM.xls
  cri_cra/cri_cra-DD-MM.csv

Suporta os padrões de nome:
  d{YY}{mon}{DD}.xls   (ex: d26mai13.xls -> debenture-13-05.xls)
  taxas_CRI_CRA_{YYYYMMDD}.csv  (ex: taxas_CRI_CRA_20260514.csv -> cri_cra-14-05.csv)
  taxas_CRI_CRA (...).csv  (lê a data do conteúdo do CSV)
  anbima *.xlsx  (lê a data da planilha)

Uso:
  python rename_anbima.py          # lista o que seria renomeado (dry-run)
  python rename_anbima.py --apply  # aplica as renomeações
"""

import os
import re
import csv
import glob
import argparse
from datetime import datetime
from pathlib import Path

BASE_DIR       = Path(__file__).parent.resolve()
DEBENTURES_DIR = BASE_DIR / "debentures"
CRI_CRA_DIR    = BASE_DIR / "cri_cra"

PT_MONTHS = {
    "jan": 1,  "fev": 2,  "mar": 3,  "abr": 4,
    "mai": 5,  "jun": 6,  "jul": 7,  "ago": 8,
    "set": 9,  "out": 10, "nov": 11, "dez": 12,
}


def _date_from_anbima_msd_name(fname: str) -> datetime | None:
    """d26mai13.xls -> datetime(2026, 5, 13)"""
    m = re.match(r"d(\d{2})([a-z]{3})(\d{2})", fname.lower())
    if not m:
        return None
    yy, mon, dd = m.group(1), m.group(2), m.group(3)
    month = PT_MONTHS.get(mon)
    if not month:
        return None
    try:
        return datetime(2000 + int(yy), month, int(dd))
    except ValueError:
        return None


def _date_from_cri_cra_name(fname: str) -> datetime | None:
    """taxas_CRI_CRA_20260514.csv -> datetime(2026, 5, 14)"""
    m = re.search(r"(\d{8})", fname)
    if not m:
        return None
    try:
        return datetime.strptime(m.group(1), "%Y%m%d")
    except ValueError:
        return None


def _date_from_cri_cra_content(filepath: Path) -> datetime | None:
    """Lê a primeira linha de dados do CSV para obter a data."""
    for enc in ("utf-8-sig", "latin-1", "cp1252"):
        try:
            with open(filepath, encoding=enc) as f:
                reader = csv.reader(f, delimiter=";")
                next(reader, None)  # skip header
                for row in reader:
                    if row and row[0].strip():
                        date_str = row[0].strip()
                        try:
                            return datetime.strptime(date_str, "%d/%m/%Y")
                        except ValueError:
                            pass
            break
        except UnicodeDecodeError:
            continue
    return None


def _date_from_xlsx_content(filepath: Path) -> datetime | None:
    """Lê a célula de data da planilha ANBIMA (linha 4, coluna B = índice 3,1)."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        for sheet_name in ["DI_SPREAD", "IPCA_SPREAD", "PREFIXADO"]:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 3:
                    val = row[1] if len(row) > 1 else None
                    if isinstance(val, datetime):
                        wb.close()
                        return val
                    break
        wb.close()
    except Exception:
        pass

    try:
        import xlrd
        wb = xlrd.open_workbook(str(filepath))
        for sheet_name in ["DI_SPREAD", "IPCA_SPREAD", "PREFIXADO"]:
            if sheet_name not in wb.sheet_names():
                continue
            ws = wb.sheet_by_name(sheet_name)
            if ws.nrows > 3:
                val = ws.cell_value(3, 1)
                if val:
                    return xlrd.xldate_as_datetime(val, wb.datemode)
    except Exception:
        pass

    return None


def plan_renames() -> list[tuple[Path, Path]]:
    """Retorna lista de (origem, destino) para renomeação."""
    renames = []
    seen_srcs: set[Path] = set()

    # --- Debêntures ---
    DEBENTURES_DIR.mkdir(exist_ok=True)
    deb_patterns = [
        DEBENTURES_DIR / "d*.xls",
        DEBENTURES_DIR / "d*.xls?",
        BASE_DIR / "d*.xls",
        BASE_DIR / "d*.xls?",
        DEBENTURES_DIR / "anbima *.xlsx",
        DEBENTURES_DIR / "anbima*.xlsx",
        BASE_DIR / "anbima data.xlsx",
    ]
    for pat in deb_patterns:
        for src in sorted(glob.glob(str(pat))):
            src = Path(src).resolve()
            if src in seen_srcs:
                continue
            seen_srcs.add(src)
            name = src.name.lower()

            # Não renomear anbima v0.xlsx — é o arquivo base histórico
            if "v0" in name:
                continue
            # Não renomear arquivos já no formato debenture-DD-MM
            if re.match(r"debenture-\d{2}-\d{2}", name):
                continue
            # Não renomear BTG
            if "119452" in name:
                continue

            dt = _date_from_anbima_msd_name(src.stem)
            if dt is None:
                print(f"  Lendo conteúdo de {src.name} para obter data...")
                dt = _date_from_xlsx_content(src)

            if dt is None:
                print(f"  [aviso] Não foi possível determinar a data de: {src.name}")
                continue

            ext = src.suffix
            dest = DEBENTURES_DIR / f"debenture-{dt.strftime('%d-%m')}{ext}"
            if src != dest:
                renames.append((src, dest))

    # --- CRI/CRA ---
    CRI_CRA_DIR.mkdir(exist_ok=True)
    cri_patterns = [
        CRI_CRA_DIR / "taxas_CRI_CRA*.csv",
        BASE_DIR / "taxas_CRI_CRA*.csv",
    ]
    for pat in cri_patterns:
        for src in sorted(glob.glob(str(pat))):
            src = Path(src).resolve()
            if src in seen_srcs:
                continue
            seen_srcs.add(src)
            name = src.name.lower()

            # Não renomear arquivos já no formato cri_cra-DD-MM
            if re.match(r"cri_cra-\d{2}-\d{2}", name):
                continue

            dt = _date_from_cri_cra_name(src.name)
            if dt is None:
                print(f"  Lendo conteúdo de {src.name} para obter data...")
                dt = _date_from_cri_cra_content(src)

            if dt is None:
                print(f"  [aviso] Não foi possível determinar a data de: {src.name}")
                continue

            dest = CRI_CRA_DIR / f"cri_cra-{dt.strftime('%d-%m')}.csv"
            if src != dest:
                renames.append((src, dest))

    return renames


def main():
    parser = argparse.ArgumentParser(
        description="Renomeia arquivos ANBIMA para o padrão debenture-DD-MM / cri_cra-DD-MM"
    )
    parser.add_argument(
        "--apply", action="store_true",
        help="Aplica as renomeações (sem --apply, apenas lista)"
    )
    args = parser.parse_args()

    renames = plan_renames()

    if not renames:
        print("Nenhum arquivo para renomear.")
        return

    print(f"\n{'=' * 55}")
    print(f"  {'ORIGEM':<35} -> DESTINO")
    print(f"{'=' * 55}")
    for src, dest in renames:
        conflict = dest.exists() and dest != src
        flag = " [CONFLITO]" if conflict else ""
        print(f"  {src.name:<35} -> {dest.name}{flag}")

    if not args.apply:
        print(f"\n{len(renames)} arquivo(s) seriam renomeados.")
        print("Use --apply para confirmar.")
        return

    ok = err = 0
    for src, dest in renames:
        if dest.exists() and dest != src:
            print(f"  [pulado] {dest.name} já existe — apague manualmente se quiser sobrescrever")
            err += 1
            continue
        try:
            # Move para a pasta destino se necessário
            dest.parent.mkdir(parents=True, exist_ok=True)
            src.rename(dest)
            print(f"  OK  {src.name}  ->  {dest.name}")
            ok += 1
        except OSError as e:
            print(f"  ERRO {src.name}: {e}")
            err += 1

    print(f"\nRenomeados: {ok}  |  Erros/pulados: {err}")


if __name__ == "__main__":
    main()
